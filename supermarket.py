import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# =========================================================
# SUPER MARKET MANAGEMENT SYSTEM
# Python + Tkinter + Excel
# =========================================================

# Excel file location
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "supermarket.xlsx")

# ---------------------------------------------------------
# CREATE EXCEL FILE
# ---------------------------------------------------------
def create_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Products"

        headers = [
            "Product ID",
            "Product Name",
            "Category",
            "Price",
            "Quantity",
            "Supplier"
        ]

        ws.append(headers)
        wb.save(EXCEL_FILE)


create_excel()


# =========================================================
# MAIN WINDOW
# =========================================================

root = tk.Tk()
root.title("Super Market Management System")
root.geometry("1000x650")
root.minsize(850, 550)

# =========================================================
# COLORS
# =========================================================

BG = "#f4f6f8"
DARK = "#263238"
WHITE = "#ffffff"
GREEN = "#2e7d32"
BLUE = "#1565c0"
RED = "#c62828"
ORANGE = "#ef6c00"

root.configure(bg=BG)


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def clear_window():
    for widget in root.winfo_children():
        widget.destroy()


def read_products():
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Products"]

    data = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(value is not None for value in row):
            data.append(row)

    wb.close()
    return data


# =========================================================
# LOGIN PAGE
# =========================================================

def show_login():

    clear_window()

    frame = tk.Frame(root, bg=BG)
    frame.pack(expand=True)

    title = tk.Label(
        frame,
        text="SUPER MARKET",
        font=("Arial", 28, "bold"),
        bg=BG,
        fg=BLUE
    )
    title.pack(pady=(20, 5))

    subtitle = tk.Label(
        frame,
        text="Management System",
        font=("Arial", 18),
        bg=BG,
        fg=DARK
    )
    subtitle.pack(pady=(0, 25))

    box = tk.Frame(
        frame,
        bg=WHITE,
        bd=2,
        relief="groove",
        padx=35,
        pady=25
    )
    box.pack()

    tk.Label(
        box,
        text="Login",
        font=("Arial", 20, "bold"),
        bg=WHITE,
        fg=DARK
    ).grid(row=0, column=0, columnspan=2, pady=(0, 20))

    tk.Label(
        box,
        text="Username:",
        font=("Arial", 12),
        bg=WHITE
    ).grid(row=1, column=0, sticky="w", pady=8)

    username_entry = tk.Entry(
        box,
        font=("Arial", 12),
        width=25
    )
    username_entry.grid(row=1, column=1, pady=8)

    tk.Label(
        box,
        text="Password:",
        font=("Arial", 12),
        bg=WHITE
    ).grid(row=2, column=0, sticky="w", pady=8)

    password_entry = tk.Entry(
        box,
        font=("Arial", 12),
        width=25,
        show="*"
    )
    password_entry.grid(row=2, column=1, pady=8)

    def login():

        username = username_entry.get().strip()
        password = password_entry.get().strip()

        if username == "" or password == "":
            messagebox.showwarning(
                "Warning",
                "Please enter username and password."
            )
            return

        if username == "admin" and password == "1234":
            messagebox.showinfo(
                "Login",
                "Login Successful!"
            )
            show_dashboard()

        else:
            messagebox.showerror(
                "Login Failed",
                "Invalid username or password."
            )

    def clear_login():
        username_entry.delete(0, tk.END)
        password_entry.delete(0, tk.END)
        username_entry.focus()

    tk.Button(
        box,
        text="LOGIN",
        command=login,
        bg=GREEN,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=12,
        cursor="hand2"
    ).grid(row=3, column=0, pady=20)

    tk.Button(
        box,
        text="CLEAR",
        command=clear_login,
        bg=ORANGE,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=12,
        cursor="hand2"
    ).grid(row=3, column=1, pady=20)

    username_entry.focus()


# =========================================================
# DASHBOARD
# =========================================================

def show_dashboard():

    clear_window()

    # Header
    header = tk.Frame(
        root,
        bg=BLUE,
        height=80
    )
    header.pack(fill="x")

    tk.Label(
        header,
        text="SUPER MARKET MANAGEMENT SYSTEM",
        font=("Arial", 22, "bold"),
        bg=BLUE,
        fg=WHITE
    ).pack(pady=20)

    # Welcome
    tk.Label(
        root,
        text="Welcome, Admin!",
        font=("Arial", 20, "bold"),
        bg=BG,
        fg=DARK
    ).pack(pady=25)

    # Button frame
    button_frame = tk.Frame(root, bg=BG)
    button_frame.pack()

    buttons = [
        ("Add Product", show_add_product, GREEN),
        ("View Products", show_view_products, BLUE),
        ("Search Product", show_search_product, ORANGE),
        ("Update Product", show_update_product, BLUE),
        ("Delete Product", show_delete_product, RED),
        ("Logout", show_login, DARK)
    ]

    row = 0
    col = 0

    for text, command, color in buttons:

        tk.Button(
            button_frame,
            text=text,
            command=command,
            bg=color,
            fg=WHITE,
            font=("Arial", 13, "bold"),
            width=18,
            height=2,
            cursor="hand2"
        ).grid(
            row=row,
            column=col,
            padx=12,
            pady=12
        )

        col += 1

        if col == 2:
            col = 0
            row += 1

    tk.Label(
        root,
        text="Data is stored in supermarket.xlsx",
        font=("Arial", 11),
        bg=BG,
        fg="gray"
    ).pack(pady=25)


# =========================================================
# ADD PRODUCT
# =========================================================

def show_add_product():

    clear_window()

    tk.Label(
        root,
        text="ADD PRODUCT",
        font=("Arial", 24, "bold"),
        bg=BG,
        fg=GREEN
    ).pack(pady=20)

    form = tk.Frame(
        root,
        bg=WHITE,
        padx=35,
        pady=20,
        bd=2,
        relief="groove"
    )
    form.pack()

    labels = [
        "Product ID",
        "Product Name",
        "Category",
        "Price",
        "Quantity",
        "Supplier"
    ]

    entries = {}

    for i, label in enumerate(labels):

        tk.Label(
            form,
            text=label + ":",
            font=("Arial", 12),
            bg=WHITE
        ).grid(
            row=i,
            column=0,
            sticky="w",
            pady=7
        )

        entry = tk.Entry(
            form,
            font=("Arial", 12),
            width=30
        )
        entry.grid(
            row=i,
            column=1,
            pady=7,
            padx=15
        )

        entries[label] = entry

    def add_product():

        product_id = entries["Product ID"].get().strip()
        name = entries["Product Name"].get().strip()
        category = entries["Category"].get().strip()
        price = entries["Price"].get().strip()
        quantity = entries["Quantity"].get().strip()
        supplier = entries["Supplier"].get().strip()

        if not all([
            product_id,
            name,
            category,
            price,
            quantity,
            supplier
        ]):
            messagebox.showwarning(
                "Warning",
                "Please fill all fields."
            )
            return

        try:
            price_value = float(price)
            quantity_value = int(quantity)

            if price_value < 0 or quantity_value < 0:
                raise ValueError

        except ValueError:
            messagebox.showerror(
                "Error",
                "Price must be a number and Quantity must be a whole number."
            )
            return

        # Check duplicate ID
        products = read_products()

        for product in products:
            if str(product[0]) == product_id:
                messagebox.showerror(
                    "Error",
                    "Product ID already exists."
                )
                return

        wb = load_workbook(EXCEL_FILE)
        ws = wb["Products"]

        ws.append([
            product_id,
            name,
            category,
            price_value,
            quantity_value,
            supplier
        ])

        wb.save(EXCEL_FILE)
        wb.close()

        messagebox.showinfo(
            "Success",
            "Product added successfully!"
        )

        clear_fields()

    def clear_fields():

        for entry in entries.values():
            entry.delete(0, tk.END)

        entries["Product ID"].focus()

    button_frame = tk.Frame(root, bg=BG)
    button_frame.pack(pady=20)

    tk.Button(
        button_frame,
        text="ADD PRODUCT",
        command=add_product,
        bg=GREEN,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).grid(row=0, column=0, padx=10)

    tk.Button(
        button_frame,
        text="CLEAR",
        command=clear_fields,
        bg=ORANGE,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).grid(row=0, column=1, padx=10)

    tk.Button(
        button_frame,
        text="BACK",
        command=show_dashboard,
        bg=DARK,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).grid(row=0, column=2, padx=10)


# =========================================================
# VIEW PRODUCTS
# =========================================================

def show_view_products():

    clear_window()

    tk.Label(
        root,
        text="VIEW PRODUCTS",
        font=("Arial", 24, "bold"),
        bg=BG,
        fg=BLUE
    ).pack(pady=15)

    frame = tk.Frame(root, bg=BG)
    frame.pack(fill="both", expand=True, padx=20, pady=10)

    columns = (
        "Product ID",
        "Product Name",
        "Category",
        "Price",
        "Quantity",
        "Supplier"
    )

    tree = ttk.Treeview(
        frame,
        columns=columns,
        show="headings"
    )

    for column in columns:

        tree.heading(
            column,
            text=column
        )

        tree.column(
            column,
            width=140,
            anchor="center"
        )

    scrollbar_y = ttk.Scrollbar(
        frame,
        orient="vertical",
        command=tree.yview
    )

    scrollbar_x = ttk.Scrollbar(
        frame,
        orient="horizontal",
        command=tree.xview
    )

    tree.configure(
        yscrollcommand=scrollbar_y.set,
        xscrollcommand=scrollbar_x.set
    )

    tree.grid(
        row=0,
        column=0,
        sticky="nsew"
    )

    scrollbar_y.grid(
        row=0,
        column=1,
        sticky="ns"
    )

    scrollbar_x.grid(
        row=1,
        column=0,
        sticky="ew"
    )

    frame.grid_rowconfigure(0, weight=1)
    frame.grid_columnconfigure(0, weight=1)

    products = read_products()

    for product in products:
        tree.insert(
            "",
            tk.END,
            values=product
        )

    tk.Button(
        root,
        text="BACK",
        command=show_dashboard,
        bg=DARK,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).pack(pady=15)


# =========================================================
# SEARCH PRODUCT
# =========================================================

def show_search_product():

    clear_window()

    tk.Label(
        root,
        text="SEARCH PRODUCT",
        font=("Arial", 24, "bold"),
        bg=BG,
        fg=ORANGE
    ).pack(pady=15)

    search_frame = tk.Frame(root, bg=BG)
    search_frame.pack(pady=10)

    tk.Label(
        search_frame,
        text="Enter Product ID or Name:",
        font=("Arial", 12),
        bg=BG
    ).pack(side="left", padx=5)

    search_entry = tk.Entry(
        search_frame,
        font=("Arial", 12),
        width=30
    )
    search_entry.pack(side="left", padx=5)

    columns = (
        "Product ID",
        "Product Name",
        "Category",
        "Price",
        "Quantity",
        "Supplier"
    )

    tree = ttk.Treeview(
        root,
        columns=columns,
        show="headings",
        height=12
    )

    for column in columns:
        tree.heading(column, text=column)
        tree.column(column, width=140)

    tree.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=10
    )

    def search():

        for item in tree.get_children():
            tree.delete(item)

        search_text = search_entry.get().strip().lower()

        if search_text == "":
            messagebox.showwarning(
                "Warning",
                "Please enter Product ID or Product Name."
            )
            return

        found = False

        for product in read_products():

            product_id = str(product[0]).lower()
            product_name = str(product[1]).lower()

            if (
                search_text in product_id
                or search_text in product_name
            ):
                tree.insert(
                    "",
                    tk.END,
                    values=product
                )
                found = True

        if not found:
            messagebox.showinfo(
                "Search",
                "Product not found."
            )

    tk.Button(
        root,
        text="SEARCH",
        command=search,
        bg=ORANGE,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).pack(pady=5)

    tk.Button(
        root,
        text="BACK",
        command=show_dashboard,
        bg=DARK,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).pack(pady=10)


# =========================================================
# UPDATE PRODUCT
# =========================================================

def show_update_product():

    clear_window()

    tk.Label(
        root,
        text="UPDATE PRODUCT",
        font=("Arial", 24, "bold"),
        bg=BG,
        fg=BLUE
    ).pack(pady=15)

    form = tk.Frame(
        root,
        bg=WHITE,
        padx=30,
        pady=15,
        bd=2,
        relief="groove"
    )
    form.pack()

    labels = [
        "Product ID",
        "Product Name",
        "Category",
        "Price",
        "Quantity",
        "Supplier"
    ]

    entries = {}

    for i, label in enumerate(labels):

        tk.Label(
            form,
            text=label + ":",
            font=("Arial", 11),
            bg=WHITE
        ).grid(
            row=i,
            column=0,
            sticky="w",
            pady=5
        )

        entry = tk.Entry(
            form,
            font=("Arial", 11),
            width=28
        )

        entry.grid(
            row=i,
            column=1,
            padx=10,
            pady=5
        )

        entries[label] = entry

    def load_product():

        product_id = entries["Product ID"].get().strip()

        if product_id == "":
            messagebox.showwarning(
                "Warning",
                "Enter Product ID first."
            )
            return

        for product in read_products():

            if str(product[0]) == product_id:

                entries["Product Name"].delete(0, tk.END)
                entries["Product Name"].insert(0, product[1])

                entries["Category"].delete(0, tk.END)
                entries["Category"].insert(0, product[2])

                entries["Price"].delete(0, tk.END)
                entries["Price"].insert(0, product[3])

                entries["Quantity"].delete(0, tk.END)
                entries["Quantity"].insert(0, product[4])

                entries["Supplier"].delete(0, tk.END)
                entries["Supplier"].insert(0, product[5])

                return

        messagebox.showerror(
            "Error",
            "Product not found."
        )

    def update_product():

        product_id = entries["Product ID"].get().strip()
        name = entries["Product Name"].get().strip()
        category = entries["Category"].get().strip()
        price = entries["Price"].get().strip()
        quantity = entries["Quantity"].get().strip()
        supplier = entries["Supplier"].get().strip()

        if not all([
            product_id,
            name,
            category,
            price,
            quantity,
            supplier
        ]):
            messagebox.showwarning(
                "Warning",
                "Please fill all fields."
            )
            return

        try:
            price_value = float(price)
            quantity_value = int(quantity)
        except ValueError:
            messagebox.showerror(
                "Error",
                "Invalid price or quantity."
            )
            return

        wb = load_workbook(EXCEL_FILE)
        ws = wb["Products"]

        found = False

        for row in ws.iter_rows(min_row=2):

            if str(row[0].value) == product_id:

                row[1].value = name
                row[2].value = category
                row[3].value = price_value
                row[4].value = quantity_value
                row[5].value = supplier

                found = True
                break

        wb.save(EXCEL_FILE)
        wb.close()

        if found:
            messagebox.showinfo(
                "Success",
                "Product updated successfully!"
            )
        else:
            messagebox.showerror(
                "Error",
                "Product ID not found."
            )

    button_frame = tk.Frame(root, bg=BG)
    button_frame.pack(pady=15)

    tk.Button(
        button_frame,
        text="LOAD",
        command=load_product,
        bg=ORANGE,
        fg=WHITE,
        font=("Arial", 11, "bold"),
        width=12
    ).grid(row=0, column=0, padx=5)

    tk.Button(
        button_frame,
        text="UPDATE",
        command=update_product,
        bg=GREEN,
        fg=WHITE,
        font=("Arial", 11, "bold"),
        width=12
    ).grid(row=0, column=1, padx=5)

    tk.Button(
        button_frame,
        text="BACK",
        command=show_dashboard,
        bg=DARK,
        fg=WHITE,
        font=("Arial", 11, "bold"),
        width=12
    ).grid(row=0, column=2, padx=5)


# =========================================================
# DELETE PRODUCT
# =========================================================

def show_delete_product():

    clear_window()

    tk.Label(
        root,
        text="DELETE PRODUCT",
        font=("Arial", 24, "bold"),
        bg=BG,
        fg=RED
    ).pack(pady=25)

    frame = tk.Frame(
        root,
        bg=WHITE,
        padx=30,
        pady=30,
        bd=2,
        relief="groove"
    )
    frame.pack()

    tk.Label(
        frame,
        text="Enter Product ID:",
        font=("Arial", 13),
        bg=WHITE
    ).pack(pady=10)

    id_entry = tk.Entry(
        frame,
        font=("Arial", 13),
        width=25
    )
    id_entry.pack(pady=10)

    def delete_product():

        product_id = id_entry.get().strip()

        if product_id == "":
            messagebox.showwarning(
                "Warning",
                "Please enter Product ID."
            )
            return

        answer = messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete this product?"
        )

        if not answer:
            return

        wb = load_workbook(EXCEL_FILE)
        ws = wb["Products"]

        found = False

        for row in range(2, ws.max_row + 1):

            if str(ws.cell(row, 1).value) == product_id:

                ws.delete_rows(row, 1)
                found = True
                break

        wb.save(EXCEL_FILE)
        wb.close()

        if found:

            messagebox.showinfo(
                "Success",
                "Product deleted successfully!"
            )

            id_entry.delete(0, tk.END)

        else:

            messagebox.showerror(
                "Error",
                "Product ID not found."
            )

    tk.Button(
        frame,
        text="DELETE PRODUCT",
        command=delete_product,
        bg=RED,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=18
    ).pack(pady=15)

    tk.Button(
        root,
        text="BACK",
        command=show_dashboard,
        bg=DARK,
        fg=WHITE,
        font=("Arial", 12, "bold"),
        width=15
    ).pack(pady=20)


# =========================================================
# START PROGRAM
# =========================================================

show_login()

root.mainloop()
